"""Estilo do manual: gera o documento de referência do pandoc (modo 'ref') e ajusta o .docx gerado (modo 'pos')."""
import sys, copy
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
AZUL_HEX = '1F3A5F'
CINZA_HEX = 'F2F4F7'
CORPO = 'Carlito'      # métrica igual à do Calibri; no Word aparece como Calibri se Carlito não existir
MONO = 'DejaVu Sans Mono'


def fonte(estilo, nome, tam=None, cor=None, negrito=None):
    f = estilo.font
    f.name = nome
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(a), nome)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if tam: f.size = Pt(tam)
    if cor is not None: f.color.rgb = cor
    if negrito is not None: f.bold = negrito


def sombra(elem_pr, cor):
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), cor)
    elem_pr.append(shd)


def campo_pagina(par):
    for tipo, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        r = OxmlElement('w:r')
        if tipo:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), tipo); r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; r.append(it)
        par._p.append(r)


def estilos(d):
    return {x.name: x for x in d.styles}


def ref(caminho):
    d = Document(caminho)
    s = estilos(d)
    for nome in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
        if nome in s:
            fonte(s[nome], CORPO, 10.5)
    for nome in ('Body Text', 'First Paragraph'):
        pf = s[nome].paragraph_format; pf.space_after = Pt(6); pf.space_before = Pt(0); pf.line_spacing = 1.1
    tam = {1: 18, 2: 14, 3: 12, 4: 11}
    for n, t in tam.items():
        st = s[f'Heading {n}']
        fonte(st, CORPO, t, AZUL, True)
        st.font.italic = False
        st.paragraph_format.space_before = Pt(18 if n == 1 else 12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    # capítulo começa em página nova, com linha embaixo
    h1 = s['Heading 1']
    h1.paragraph_format.page_break_before = True
    ppr = h1.element.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', '8'), ('w:space', '4'), ('w:color', AZUL_HEX)): b.set(qn(k), v)
    bdr.append(b); ppr.append(bdr)
    fonte(s['Title'], CORPO, 30, AZUL, True)
    s['Title'].paragraph_format.space_before = Pt(160)
    fonte(s['Subtitle'], CORPO, 15, RGBColor(0x44, 0x54, 0x6A), False)
    for nome in ('Author', 'Date'):
        fonte(s[nome], CORPO, 12, RGBColor(0x44, 0x54, 0x6A))
    fonte(s['Verbatim Char'], MONO, 8)
    # citações (notas "> ...")
    if 'Block Text' in s:
        bt = s['Block Text']; fonte(bt, CORPO, 10, RGBColor(0x33, 0x33, 0x33))
        bppr = bt.element.get_or_add_pPr(); sombra(bppr, 'FFF6E0')
        bt.paragraph_format.left_indent = Cm(0.4)
    # página A4, margens 2 cm, rodapé com número
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(2)
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        p.text = 'Manual de Implantação Scout + eLux — POPULOS    |    página '
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in p.runs: r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        campo_pagina(p)
    d.save(caminho)


def bordas_tabela(tbl):
    tblpr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{lado}')
        for k, v in (('w:val', 'single'), ('w:sz', '4'), ('w:space', '0'), ('w:color', 'B7C3D0')): e.set(qn(k), v)
        b.append(e)
    velho = tblpr.find(qn('w:tblBorders'))
    if velho is not None: tblpr.remove(velho)
    tblpr.append(b)
    # largura 100% da página
    w = tblpr.find(qn('w:tblW'))
    if w is None: w = OxmlElement('w:tblW'); tblpr.append(w)
    w.set(qn('w:type'), 'pct'); w.set(qn('w:w'), '5000')
    mar = OxmlElement('w:tblCellMar')
    for lado in ('left', 'right'):
        e = OxmlElement(f'w:{lado}'); e.set(qn('w:w'), '80'); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tblpr.append(mar)


def pos(caminho):
    d = Document(caminho)
    s = estilos(d)
    for nome in ('Source Code', 'Verbatim Char'):
        if nome in s:
            fonte(s[nome], MONO, 8)
    if 'Source Code' in s:
        sc = s['Source Code']
        sc.paragraph_format.space_before = Pt(4); sc.paragraph_format.space_after = Pt(6)
        sombra(sc.element.get_or_add_pPr(), CINZA_HEX)
    for tbl in d.tables:
        bordas_tabela(tbl)
        for i, row in enumerate(tbl.rows):
            for cell in row.cells:
                tcpr = cell._tc.get_or_add_tcPr()
                if i == 0:
                    sombra(tcpr, AZUL_HEX)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
                    for r in p.runs:
                        rs = r._r.rPr.find(qn('w:rStyle')) if r._r.rPr is not None else None
                        r.font.size = Pt(7.5) if rs is not None and rs.get(qn('w:val')) == 'VerbatimChar' else Pt(8.5)
                        if i == 0:
                            r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # cabeçalho repete em cada página
        trpr = tbl.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trpr.append(th)
    # o Word atualiza o sumário (números de página) ao abrir
    uf = OxmlElement('w:updateFields'); uf.set(qn('w:val'), 'true')
    d.settings.element.append(uf)
    d.save(caminho)


if __name__ == '__main__':
    {'ref': ref, 'pos': pos}[sys.argv[1]](sys.argv[2])
